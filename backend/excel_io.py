"""Load .xlsx / .xls workbooks into openpyxl Workbook objects."""
from __future__ import annotations

import io
import re
from struct import error as StructError
from struct import unpack

import openpyxl
from openpyxl.workbook import Workbook


XLSX_MAGIC = b"PK"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"  # OLE Compound Document
_ILLEGAL_CELL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BAD_TITLE_RE = re.compile(r'[:\\/?*\[\]]')
_XLRD_SST_PATCHED = False


def sniff_excel_kind(content: bytes, filename: str | None = None) -> str:
    name = (filename or "").lower()
    if content.startswith(XLSX_MAGIC) or name.endswith(".xlsx"):
        return "xlsx"
    if content.startswith(XLS_MAGIC) or name.endswith(".xls"):
        return "xls"
    raise ValueError("Ожидается файл Excel (.xlsx или .xls)")


def load_workbook(content: bytes, *, data_only: bool = True, filename: str | None = None) -> Workbook:
    kind = sniff_excel_kind(content, filename)
    label = filename or kind
    try:
        if kind == "xlsx":
            return openpyxl.load_workbook(io.BytesIO(content), data_only=data_only)
        return _load_xls_as_workbook(content)
    except AssertionError as e:
        raise RuntimeError(
            f"Не удалось прочитать «{label}»: файл Excel повреждён или в неподдерживаемом формате"
        ) from e


def _safe_sheet_title(name: str | None, used: set[str]) -> str:
    title = _BAD_TITLE_RE.sub("_", (name or "Sheet")).strip()[:31] or "Sheet"
    base = title
    n = 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = (base[: 31 - len(suffix)] + suffix) or f"Sheet{n}"
        n += 1
    used.add(title.lower())
    return title


def _safe_cell_value(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    if not isinstance(value, str):
        try:
            value = str(value)
        except Exception:
            return None
    value = _ILLEGAL_CELL_RE.sub("", value)
    if not value:
        return None
    return value[:32767]


def _unpack_sst_table_tolerant(datatab, nstrings):
    """xlrd SST parser that tolerates CONTINUE-record boundaries Excel itself accepts.

    Some .xls files (often Save-As from xlsx / Asian exporters) declare a shared
    string table whose last strings start a few bytes before a record split.
    Stock xlrd then hits ``assert _unused_i == nstrings - 1``.
    """
    from xlrd.timemachine import BYTES_ORD, UNICODE_LITERAL, xrange

    datainx = 0
    ndatas = len(datatab)
    data = datatab[0]
    datalen = len(data)
    pos = 8
    strings = []
    richtext_runs = {}
    latin_1 = "latin_1"

    def ensure(n: int) -> bool:
        nonlocal datainx, data, datalen, pos
        while pos + n > datalen:
            if datainx + 1 >= ndatas:
                return False
            leftover = data[pos:] if pos < datalen else b""
            datainx += 1
            data = leftover + datatab[datainx]
            datalen = len(data)
            pos = 0
        return True

    for _i in xrange(nstrings):
        try:
            if not ensure(3):
                break
            nchars = unpack("<H", data[pos:pos + 2])[0]
            pos += 2
            options = BYTES_ORD(data[pos])
            pos += 1
            rtcount = 0
            phosz = 0
            if options & 0x08:
                if not ensure(2):
                    break
                rtcount = unpack("<H", data[pos:pos + 2])[0]
                pos += 2
            if options & 0x04:
                if not ensure(4):
                    break
                phosz = unpack("<i", data[pos:pos + 4])[0]
                pos += 4
            accstrg = UNICODE_LITERAL("")
            charsgot = 0
            while True:
                charsneed = nchars - charsgot
                if options & 0x01:
                    charsavail = min((datalen - pos) >> 1, charsneed)
                    accstrg += str(data[pos:pos + 2 * charsavail], "utf_16_le")
                    pos += 2 * charsavail
                else:
                    charsavail = min(datalen - pos, charsneed)
                    accstrg += str(data[pos:pos + charsavail], latin_1)
                    pos += charsavail
                charsgot += charsavail
                if charsgot >= nchars:
                    break
                datainx += 1
                if datainx >= ndatas:
                    break
                data = datatab[datainx]
                datalen = len(data)
                options = BYTES_ORD(data[0])
                pos = 1
            extra = rtcount * 4 + max(phosz, 0)
            while extra:
                if not ensure(1):
                    break
                take = min(extra, datalen - pos)
                pos += take
                extra -= take
            strings.append(accstrg)
        except (AssertionError, IndexError, StructError, UnicodeDecodeError):
            break
    return strings, richtext_runs


def _patch_xlrd_sst() -> None:
    global _XLRD_SST_PATCHED
    if _XLRD_SST_PATCHED:
        return
    import xlrd.book as bookmod

    bookmod.unpack_SST_table = _unpack_sst_table_tolerant
    _XLRD_SST_PATCHED = True


def _load_xls_as_workbook(content: bytes) -> Workbook:
    try:
        import xlrd
    except ImportError as e:
        raise RuntimeError(
            "Для .xls установите зависимость xlrd: pip install xlrd==1.2.0"
        ) from e

    _patch_xlrd_sst()
    try:
        book = xlrd.open_workbook(file_contents=content, logfile=io.StringIO())
    except AssertionError as e:
        raise RuntimeError(
            "Не удалось прочитать .xls: файл в старом или повреждённом формате Excel"
        ) from e

    out = Workbook()
    out.remove(out.active)
    used_titles: set[str] = set()

    for sheet in book.sheets():
        ws = out.create_sheet(title=_safe_sheet_title(sheet.name, used_titles))
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(cell.value)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = None
                try:
                    ws.cell(row=r + 1, column=c + 1, value=_safe_cell_value(value))
                except Exception:
                    continue
    if not out.worksheets:
        out.create_sheet("Sheet1")
    return out
