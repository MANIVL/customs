"""Build a catalog-shaped table from article codes found in an arbitrary Excel file."""
from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Font

import articles
import countries
import excel_io

# Headers that mark an article column even when none of its codes are in the catalog yet.
_HEADER_EXACT = {
    "артикул",
    "артикул товара",
    "код артикула",
    "article",
    "article no",
    "article number",
    "sku",
    "item",
    "item no",
    "item number",
    "part no",
    "part number",
    "product code",
}
_HEADER_SCAN_ROWS = 25


def _token(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip().upper()
    if isinstance(value, int):
        return str(value)
    return " ".join(str(value).strip().upper().split())


def _is_article_header(text: str) -> bool:
    norm = " ".join(text.casefold().split())
    if not norm:
        return False
    if norm in _HEADER_EXACT or norm.startswith("артикул"):
        return True
    return False


def _grids(content: bytes, filename: str | None) -> list[tuple[str, list[list[str]]]]:
    book = excel_io.load_workbook(content, data_only=True, filename=filename)
    grids: list[tuple[str, list[list[str]]]] = []
    for sheet in book.worksheets:
        rows = [[_token(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        grids.append((sheet.title, rows))
    return grids


def _looks_like_article(token: str) -> bool:
    if _is_article_header(token) or len(token) < 3:
        return False
    return any(ch.isdigit() for ch in token) or "-" in token or "/" in token


def _column_values(rows: list[list[str]], col: int, start: int, index: dict[str, dict[str, Any]]) -> list[str]:
    values: list[str] = []
    for row in rows[start:]:
        if col >= len(row):
            continue
        token = row[col]
        if not token or _is_article_header(token):
            continue
        if token.casefold() in index or _looks_like_article(token):
            values.append(token)
    return values


def _best_column(
    grids: list[tuple[str, list[list[str]]]],
    index: dict[str, dict[str, Any]],
) -> tuple[int, int, int, bool] | None:
    """Return (sheet, column, data start, header was recognized) for the article column."""
    header_choices: list[tuple[int, int, int, int, int]] = []
    plain_choices: list[tuple[int, int, int, int]] = []
    for sheet_i, (_title, rows) in enumerate(grids):
        width = max((len(row) for row in rows), default=0)
        header_at: dict[int, int] = {}
        for row_i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
            for col, token in enumerate(row):
                if col not in header_at and _is_article_header(token):
                    header_at[col] = row_i
        for col in range(width):
            start = header_at[col] + 1 if col in header_at else 0
            values = _column_values(rows, col, start, index)
            hits = sum(1 for value in values if value.casefold() in index)
            if col in header_at:
                header_choices.append((hits, len(values), sheet_i, col, start))
            elif hits:
                plain_choices.append((hits, sheet_i, col, hits))
    if header_choices:
        header_choices.sort(key=lambda item: (item[0], item[1]), reverse=True)
        _hits, _count, sheet_i, col, start = header_choices[0]
        return sheet_i, col, start, True
    if not plain_choices:
        return None
    plain_choices.sort(reverse=True)
    best_hits, sheet_i, col, _hits = plain_choices[0]
    second = plain_choices[1][0] if len(plain_choices) > 1 else 0
    if best_hits < 1 or best_hits == second:
        return None
    return sheet_i, col, 0, False


_COMPARE_KEYS = ("hs_code", "origin_code", "group_description")


def _compare_kind(token: str) -> str | None:
    norm = " ".join(token.casefold().split())
    if not norm or "описание товара" in norm:
        return None
    if "описание группы" in norm or norm == "описание":
        return "group_description"
    if "происхожд" in norm or norm in {"origin", "origin code", "country", "country of origin"}:
        return "origin_code"
    if "код страны" in norm:
        return "origin_code"
    if (
        "код товара" in norm
        or "тн вэд" in norm
        or "tn ved" in norm
        or norm in {"hs code", "hs", "коды", "hscode", "hs-code"}
    ):
        return "hs_code"
    if "hs" in norm and "code" in norm:
        return "hs_code"
    return None


def _ten_digit_count(rows: list[list[str]], col: int) -> int:
    return sum(1 for row in rows if col < len(row) and _full_hs(row[col]))


def _compare_map(header: list[str], data_rows: list[list[str]], *, keep_hs_header: bool = False) -> dict[str, int]:
    mapping: dict[str, int] = {}
    hs_candidates: list[int] = []
    for col, token in enumerate(header):
        kind = _compare_kind(token)
        if kind == "hs_code":
            hs_candidates.append(col)
            continue
        if kind and kind not in mapping:
            mapping[kind] = col
    if hs_candidates:
        best = max(hs_candidates, key=lambda col: _ten_digit_count(data_rows, col))
        if keep_hs_header or _ten_digit_count(data_rows, best):
            mapping["hs_code"] = best
    return mapping


def _row_gaps(row: list[str], compared: dict[str, int]) -> list[str]:
    gaps: list[str] = []
    desc_col = compared.get("group_description")
    hs_col = compared.get("hs_code")
    description = row[desc_col] if desc_col is not None and desc_col < len(row) else ""
    hs_code = row[hs_col] if hs_col is not None and hs_col < len(row) else ""
    if not description:
        gaps.append("описание")
    if not _full_hs(hs_code):
        gaps.append("код ТН ВЭД")
    return gaps


def _blank_item(code: str) -> dict[str, str]:
    item = {key: "" for key in articles.FIELD_KEYS}
    item["article"] = code
    return item


def _catalog_item(row_data: dict[str, Any]) -> dict[str, str]:
    return {key: row_data[key] for key in articles.FIELD_KEYS}


def _full_hs(value: str) -> bool:
    """A product code is comparable only in the full 10-digit form."""
    return len(value) == 10 and value.isdigit()


def _values_differ(key: str, file_value: str, db_value: str) -> bool:
    if not file_value or file_value == db_value:
        return False
    if key == "hs_code":
        return _full_hs(file_value) and _full_hs(db_value)
    return True


def build_load_table(content: bytes, filename: str | None = None, *, compare: bool = False) -> dict[str, Any]:
    if not content:
        raise ValueError("Файл пустой")
    grids = _grids(content, filename)
    if not grids or not any(rows for _title, rows in grids):
        raise ValueError("В файле нет данных")
    index = articles.index_by_article()
    chosen = _best_column(grids, index)
    items: list[dict[str, str]] = []
    diffs: list[dict[str, Any]] = []
    seen_diffs: set[str] = set()
    compared: dict[str, int] = {}
    if compare and (chosen is None or not chosen[3]):
        raise ValueError("Для сравнения в файле нужен столбец «Артикул». Заполните шаблон.")
    if chosen is None:
        mode = "scan"
        for _title, rows in grids:
            for row in rows:
                for token in row:
                    if not token or _is_article_header(token):
                        continue
                    row_data = index.get(token.casefold())
                    if row_data is None:
                        continue
                    items.append(_catalog_item(row_data))
    else:
        mode = "column"
        sheet_i, col, start, _header = chosen
        rows = grids[sheet_i][1]
        compared = _compare_map(
            rows[start - 1] if start else [],
            rows[start:],
            keep_hs_header=compare,
        )
        if compare and ("group_description" not in compared or "hs_code" not in compared):
            raise ValueError(
                "Для сравнения нужны столбцы «Описание» и «Код ТН ВЭД». Скачайте шаблон и заполните их."
            )
        if compare:
            incomplete: list[str] = []
            for row in rows[start:]:
                code = row[col] if col < len(row) else ""
                if not code or _is_article_header(code) or code.casefold() not in index:
                    continue
                gaps = _row_gaps(row, compared)
                if gaps:
                    incomplete.append(f"{code} ({', '.join(gaps)})")
            if incomplete:
                shown = ", ".join(incomplete[:8])
                extra = len(incomplete) - 8
                if extra > 0:
                    shown += f" и ещё {extra}"
                raise ValueError(
                    "Для сравнения заполните описание и код ТН ВЭД из 10 цифр. "
                    "Страну происхождения можно оставить пустой. Не заполнено: " + shown
                )
        for row in rows[start:]:
            code = row[col] if col < len(row) else ""
            if not code or _is_article_header(code):
                continue
            if code.casefold() not in index and not _looks_like_article(code):
                continue
            row_data = index.get(code.casefold())
            items.append(_catalog_item(row_data) if row_data else _blank_item(code))
            if code.casefold() in seen_diffs:
                continue
            file_values = {}
            for key, pos in compared.items():
                value = row[pos] if pos < len(row) else ""
                if row_data is not None and key == "hs_code" and not _full_hs(value):
                    value = ""
                if key == "origin_code" and value:
                    value = countries.resolve_country(value) or value
                file_values[key] = value
            file_compared = {key: file_values.get(key, "") for key in _COMPARE_KEYS}
            if row_data is None:
                seen_diffs.add(code.casefold())
                diffs.append(
                    {
                        "id": None,
                        "article": code,
                        "missing": True,
                        "db": _blank_item(code),
                        "file": file_compared,
                        "fields": ["missing"],
                    }
                )
                continue
            changed = [
                key
                for key in _COMPARE_KEYS
                if _values_differ(key, file_compared.get(key, ""), row_data[key] or "")
            ]
            if not changed:
                continue
            seen_diffs.add(code.casefold())
            diffs.append(
                {
                    "id": row_data["id"],
                    "article": row_data["article"],
                    "missing": False,
                    "db": _catalog_item(row_data),
                    "file": file_compared,
                    "fields": changed,
                }
            )
    found_count = sum(1 for item in items if any(item[key] for key in articles.FIELD_KEYS if key != "article"))
    return {
        "mode": mode,
        "sheet": grids[chosen[0]][0] if chosen else "",
        "total": len(items),
        "found_count": found_count,
        "items": items,
        "fields": articles.field_meta(),
        "diffs": diffs,
        "compared_fields": list(compared),
    }


def load_table_xlsx(result: dict[str, Any]) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Загрузочная таблица"
    fields = result.get("fields") or articles.field_meta()
    headers = [field["label"] for field in fields]
    keys = [field["key"] for field in fields]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for item in result.get("items") or []:
        sheet.append([item.get(key) or "" for key in keys])
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(12, min(48, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
